from datetime import datetime
from pathlib import Path
from importlib.util import spec_from_file_location,module_from_spec
ROOT=Path(__file__).resolve().parents[1]
def ok(c): return isinstance(c,dict) and c.get('role')=='MEMBER' and bool(c.get('user_id'))
def db():
 s=spec_from_file_location('db',ROOT/'config/member_database.py');m=module_from_spec(s);s.loader.exec_module(m);return m.get_database()
def _member(c):
 try:
  client,d=db();x=d.members.find_one({'user_id':c['user_id']},{'_id':0,'name':1,'mobile':1});client.close();return x or {}
 except Exception:return {}
def list_months(c,month):
 if not ok(c): raise PermissionError('Access denied')
 try:
  client,d=db();rows=list(d.member_monthly_records.find({'user_id':c['user_id'],'month':month},{'_id':0}).limit(1));client.close();return rows
 except Exception:return []
def upsert_month(c,data):
 if not ok(c): raise PermissionError('Access denied')
 month=str(data.get('month') or datetime.now().strftime('%Y-%m'));m=_member(c)
 x={'user_id':c['user_id'],'name':data.get('name',m.get('name',c.get('name',''))),'mobile':data.get('mobile',m.get('mobile',c.get('mobile',''))),'meal_type':data.get('meal_type',''),'month':month,'date':data.get('date',datetime.now().strftime('%Y-%m-%d')),'description':data.get('description',''),'updated_at':datetime.utcnow().isoformat()}
 client,d=db();d.member_monthly_records.update_one({'user_id':c['user_id'],'month':month},{'$set':x},upsert=True);d.member_monthly_records.create_index([('user_id',1),('month',1)],unique=True);row=d.member_monthly_records.find_one({'user_id':c['user_id'],'month':month},{'_id':0});client.close();return row
def delete_month(c,month):
 if not ok(c): raise PermissionError('Access denied')
 client,d=db();r=d.member_monthly_records.delete_one({'user_id':c['user_id'],'month':month});client.close();return r.deleted_count==1
def export_doc(c,month):
 if not ok(c): raise PermissionError('Access denied')
 rows=list_months(c,month)
 row=rows[0] if rows else {'name':c.get('name',''),'mobile':c.get('mobile',''),'meal_type':'','month':month,'date':'','description':''}
 from docx import Document
 from io import BytesIO
 doc=Document();doc.add_heading('MEMBER MONTHLY RECORD',0)
 table=doc.add_table(rows=1,cols=8)
 for i,v in enumerate(['SL','Name','Mobile','Meal type','Month','Date','Description','Action']): table.rows[0].cells[i].text=v
 vals=['1',str(row.get('name','')),str(row.get('mobile','')),str(row.get('meal_type','')),str(row.get('month','')),str(row.get('date','')),str(row.get('description','')),'Update / Download']
 for i,v in enumerate(vals): table.add_row().cells[i].text=v
 out=BytesIO();doc.save(out);out.seek(0);return out
